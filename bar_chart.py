from openpyxl import load_workbook
import matplotlib.pyplot as plt

# Load the Excel file
workbook = load_workbook("data.xlsx", data_only=True)
sheet = workbook.active

# Read data from the first two columns
categories = []
values = []

for row in sheet.iter_rows(min_row=2, values_only=True):
    categories.append(row[0])
    values.append(row[1])

# Create the bar chart
plt.bar(categories, values, color="skyblue")

# Add labels and title
plt.xlabel("Category")
plt.ylabel("Value")
plt.title("Bar Chart")

# Rotate category names if needed
plt.xticks(rotation=45)

# Save and display the chart
plt.tight_layout()
plt.savefig("bar_chart.png")
plt.show()
